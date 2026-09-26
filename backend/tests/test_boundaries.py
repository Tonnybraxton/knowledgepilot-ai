import io
import uuid
import zipfile

import pytest
from docx import Document
from pydantic import ValidationError
from reportlab.pdfgen.canvas import Canvas

from app.ai.chunking import chunk_passages
from app.ai.context import build_context, retrieval_query, validate_citations
from app.ai.extraction import Passage, extract
from app.ai.retrieval import reciprocal_rank_fusion
from app.core.config import Settings, settings
from app.core.errors import AppError
from app.schemas.api import RegisterIn, Scope, SearchHit
from app.services.uploads import MIMES, validate_upload


@pytest.mark.parametrize(
    "name,content",
    [
        ("script.exe", b"MZ"),
        ("report.pdf", b"text"),
        ("empty.txt", b""),
        ("bad.txt", b"a\x00b"),
        ("bad.md", b"\xff"),
        ("bad.docx", b"not zip"),
    ],
)
def test_rejects_unsafe_uploads(name: str, content: bytes) -> None:
    with pytest.raises(AppError):
        validate_upload(name, content)


def test_rejects_macro_docx() -> None:
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        for name in ["word/document.xml", "[Content_Types].xml", "word/vbaProject.bin"]:
            archive.writestr(name, "content")
    with pytest.raises(AppError):
        validate_upload("macro.docx", data.getvalue())


def test_upload_size_and_filename(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings(), "max_upload_bytes", 10)
    assert validate_upload("../notes.txt", b"hello")[0] == ".._notes.txt"
    with pytest.raises(AppError):
        validate_upload("notes.txt", b"x" * 11)


def test_pdf_page_provenance() -> None:
    data = io.BytesIO()
    pdf = Canvas(data)
    for text in ["Orion uses cobalt batteries.", "Orion launches in October."]:
        pdf.drawString(72, 700, text)
        pdf.showPage()
    pdf.save()
    passages = extract(data.getvalue(), "application/pdf")
    assert [passage.page for passage in passages] == [1, 2]
    assert "October" in passages[1].text


def test_docx_sections_and_tables() -> None:
    document = Document()
    document.add_heading("Specifications", level=1)
    document.add_paragraph("Orion uses cobalt batteries.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Capacity 80 kWh"
    data = io.BytesIO()
    document.save(data)
    passages = extract(data.getvalue(), MIMES[".docx"])
    assert any(
        "80 kWh" in passage.text and passage.section == "Specifications" for passage in passages
    )


def test_empty_extraction_abstains() -> None:
    with pytest.raises(ValueError, match="No readable text"):
        extract(b"   ", "text/plain")


def test_chunk_bounds_overlap_and_repeated_page_provenance(monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = settings()
    monkeypatch.setattr(cfg, "chunk_target_tokens", 30)
    monkeypatch.setattr(cfg, "chunk_max_tokens", 40)
    monkeypatch.setattr(cfg, "chunk_overlap_tokens", 5)
    chunks = chunk_passages(
        [
            Passage("A repeated header.", 1),
            Passage("A repeated header.", 2),
            Passage("Orion cobalt battery specifications. " * 100, 3),
        ]
    )
    assert [chunk.page_number for chunk in chunks[:2]] == [1, 2]
    assert all(0 < chunk.token_count <= 40 for chunk in chunks)
    assert len(chunks) > 3


def test_citation_validation_and_context_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    assert validate_citations("A [1], invented [9], zero [0].", 2) == (
        "A [1], invented , zero .",
        {1},
    )
    hit = SearchHit(
        chunk_id=uuid.uuid4(),
        document_id=uuid.uuid4(),
        document_name="notes.md",
        content="Evidence from the document",
        page_number=3,
        section_title=None,
        score=1,
    )
    context, sources = build_context([hit])
    assert '"citation": 1' in context and sources == [hit]
    monkeypatch.setattr(settings(), "context_max_tokens", 1)
    assert build_context([hit]) == ("", [])


def test_rank_fusion_deduplicates_each_ranking() -> None:
    a, b = uuid.uuid4(), uuid.uuid4()
    scores = reciprocal_rank_fusion([[a, a, b], [b, a]])
    assert scores[a] == scores[b]


def test_followup_is_bounded() -> None:
    assert "Follow-up" in retrieval_query("What about its cost?", ["x" * 5000])
    assert len(retrieval_query("Compare those", ["x" * 5000])) < 1600


def test_configuration_and_input_invariants() -> None:
    with pytest.raises(ValidationError):
        Settings(embedding_dimensions=768)
    with pytest.raises(ValidationError):
        Settings(app_env="production", frontend_url="http://localhost")
    with pytest.raises(ValidationError):
        RegisterIn(email="person@example.com", password="long-password", display_name="  ")
    with pytest.raises(ValidationError):
        Scope(document_ids=[uuid.uuid4() for _ in range(21)])
